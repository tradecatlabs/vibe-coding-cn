#!/usr/bin/env python3
"""离线维护同一功法总表的Markdown/Excel视图；默认只检查，不评级或写登记。

运行：make sync-gongfa-catalog / make check-gongfa-catalog。
依赖：requirements-gongfa.txt；复用登记校验器与openpyxl，不访问临时分析目录。
"""
import argparse
from collections import Counter, defaultdict
from hashlib import sha256
from html import escape
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import re
import sys
import tempfile
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('gongfa_protocol', Path(__file__).with_name('check-gongfa.py'))
protocol = importlib.util.module_from_spec(spec)
spec.loader.exec_module(protocol)
MARKER = '<!-- generated: sync-gongfa-catalog.py -->'
CREATOR = 'vibe-coding-cn:sync-gongfa-catalog.py'
GRADES = [f'{tier}-{level}' for tier in ('huang', 'xuan', 'di', 'tian')
          for level in ('initial', 'intermediate', 'advanced')]
GRADE_LABELS = {f'{tier}-{level}': a + b for tier, a in [('huang', '黄'), ('xuan', '玄'), ('di', '地'), ('tian', '天')]
                for level, b in [('initial', '初'), ('intermediate', '中'), ('advanced', '高')]}
TYPE_LABELS = {'explanatory-content': '解释型', 'normative-allocation-content': '准则型',
               'method-specification': '方法型', 'dao-content': '复合（保留上位）', None: '待分类'}
BATCH_LABELS = {'registry': '原登记库', 'html211': '此前四HTML', 'solve24': 'Solve已审首24', 'html96': 'HTML续提'}
HEADERS = ['序号', '暂定品级', '功法名称', '本体类别', '登记状态', '来源批次', '用途/任务族',
           '限定内容（原文或源条目）', '边界/限制', '初评理由', '反证条件', '主观置信', '原始ID',
           '内容版本', '来源与选区', '限定内容SHA-256', '关联/重名提示', '原初评时间']


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(text):
    return sha256(text.encode('utf-8')).hexdigest()


def as_text(value):
    return '' if value is None else value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)


def unique(items, key):
    result = {}
    for item in items:
        identity = key(item)
        require(identity not in result, '来源对象、内容版本或初评意见重复')
        result[identity] = item
    return result


def load_inputs(root):
    root = root.resolve(strict=True)
    registry_path = protocol.safe_path(root, 'metadata/gongfa/registry.json')
    registry_sha = sha256(protocol.bounded_read(registry_path, 32 * 1024 * 1024)).hexdigest()
    registry, registry_summary = protocol.validate(root, registry_path)
    require(sha256(protocol.bounded_read(registry_path, 32 * 1024 * 1024)).hexdigest() == registry_sha, 'registry.json读取期间变化，请重跑')
    index_path = protocol.safe_path(root, 'metadata/gongfa/catalog-sources.json')
    index = protocol.load_json(index_path, 1024 * 1024)
    require(set(index) == {'schema_version', 'kind', 'snapshot', 'candidate_counts'} and index['schema_version'] == '1.0.0'
            and index['kind'] == 'catalog_candidate_source_index_not_registry', 'catalog-sources.json协议形态不符')
    reference = index['snapshot']
    require(set(reference) == {'path', 'sha256'} and re.fullmatch(r'source-[a-z0-9-]+\.json', reference['path']), '候选快照须为本目录source-*.json')
    require(not (root / 'metadata/gongfa' / reference['path']).is_symlink(), '候选快照不能是软链接')
    snapshot_path = protocol.safe_path(root, 'metadata/gongfa/' + reference['path'])
    raw = protocol.bounded_read(snapshot_path, 32 * 1024 * 1024)
    require(sha256(raw).hexdigest() == reference['sha256'], '候选来源快照SHA-256不符')
    snapshot = protocol.load_json(snapshot_path)
    require(sha256(protocol.bounded_read(snapshot_path, 32 * 1024 * 1024)).hexdigest() == reference['sha256'], '候选快照读取期间变化')
    require(set(snapshot) == {'schema_version', 'kind', 'captured_at', 'scope', 'source_input_bindings',
                            'independent_semantic_approval', 'source_books_and_licenses_verified', 'formal_assessments_created', 'documents'}
            and snapshot['schema_version'] == '1.0.0' and snapshot['kind'] == 'frozen_classified_candidate_source_not_registry', '候选来源快照协议形态不符')
    require(snapshot['independent_semantic_approval'] is False and snapshot['source_books_and_licenses_verified'] is False
            and snapshot['formal_assessments_created'] == 0, '历史来源快照不能冒充独立批准、原著核验或正式评级')
    protocol.timestamp(snapshot['captured_at'], 'candidate_snapshot/captured_at')
    documents = snapshot['documents']
    require(set(documents) == {'html211_candidates', 'html211_grades', 'html96_candidates', 'html96_grades', 'solve24_analysis', 'solve_raw'}, '候选来源文档集合不符')
    counts = {'html211': len(documents['html211_candidates']['candidates']), 'html96': len(documents['html96_candidates']['candidates']),
              'solve24': len(documents['solve24_analysis']['items'])}
    require(counts == index['candidate_counts'], '候选批次范围与索引不一致')
    reviewed = {item['source_id'] for item in documents['solve24_analysis']['items']}
    require({item['source_id'] for item in documents['solve_raw']['items']} == reviewed, 'Solve原条目须仅含已审范围')
    binding = {'registry_sha256': registry_sha, 'candidate_sha256': reference['sha256'], 'candidate_snapshot': reference['path'],
               'source_captured_at': snapshot['captured_at'], 'registry_assessments': registry_summary['assessments']}
    return {**documents, 'registry': registry}, binding


def merge_rows(documents):
    rows = []
    registry = documents['registry']
    sources = unique(registry['sources'], lambda item: item['id'])
    batches = unique(registry['grade_proposals'], lambda batch: batch['id'])
    links = {batch['id']: batch['supersedes'] for batch in batches.values()}
    protocol.check_chains(links, 'grade_proposals/supersedes', reject_forks=True)
    superseded = {value for value in links.values() if value is not None}
    opinions = defaultdict(list)
    for batch in batches.values():
        if batch['id'] not in superseded:
            unique(batch['items'], lambda item: (item['gongfa_id'], item['content_version']))
            for item in batch['items']:
                opinions[item['gongfa_id'], item['content_version']].append((item, batch))
    entries = unique(registry['gongfa'], lambda item: (item['id'], item['content_version']))
    for item in entries.values():
        content = item['content']
        require(content['text_sha256'] == digest(content['original_text']), '登记正文SHA-256不符')
        proposals = opinions[item['id'], item['content_version']]
        for grade, _batch in proposals:
            require(grade['content_sha256'] == content['text_sha256'], '登记初评正文绑定不符')
        grade = proposals[0][0] if len(proposals) == 1 else {}
        batch = proposals[0][1] if len(proposals) == 1 else {}
        status = 'provisional' if len(proposals) == 1 else 'multiple-contexts' if proposals else 'unrated'
        detail = as_text([{'batch_id': b['id'], 'application_context': b['application_context'],
                           'baseline_assumption': b['baseline_assumption'], 'assessed_at': b['assessed_at'], 'item': g} for g, b in proposals])
        reference = item['source_ref']
        boundary = '选区说明：' + content['selection_note']
        for field, label in [('conditions', '条件'), ('limitations', '限制')]:
            if content['profile'][field]:
                boundary += '\n' + label + '：' + as_text(content['profile'][field])
        rows.append({'namespace': 'registry', 'original_id': item['id'], 'name': item['name'], 'ontology_type': item['ontology_type'],
                     'grade_id': grade.get('grade_id'), 'grading_status': status,
                     'entry_status': '已登记' if item['entry_status'] == 'registered' else '登记库候选',
                     'task_family': grade.get('task_family', ''), 'original_text': content['original_text'], 'boundary': boundary,
                     'reason': grade.get('reason', detail if proposals else '当前内容版本未初评；不继承旧版本意见。'),
                     'falsifier': grade.get('falsifier', ''), 'confidence': grade.get('confidence', ''), 'content_version': item['content_version'],
                     'source': as_text({'source': sources[reference['source_id']], 'selection': reference, 'original_proposal_context': detail}),
                     'content_sha256': content['text_sha256'], 'relations': '', 'assessed_at': batch.get('assessed_at', '')})
    for namespace in ('html211', 'html96'):
        candidates = unique(documents[namespace + '_candidates']['candidates'], lambda item: item['candidate_id'])
        grade_document = documents[namespace + '_grades']
        grades = unique(grade_document['items'], lambda item: item['candidate_id'])
        require(set(candidates) == set(grades), 'HTML候选与原初评范围不符')
        for item in candidates.values():
            grade = grades[item['candidate_id']]
            require(grade['content_sha256'] == item['text_sha256'] == digest(item['original_text']), 'HTML限定正文绑定不符')
            require(grade['source_selection'] == item['source_selection'] and grade['ontology_type'] == item['ontology_type']
                    and grade['name'] == item['name'], 'HTML选区、类型或名称与原意见不符')
            require(grade['grading_status'] in {'provisional', 'undetermined'}
                    and (grade['provisional_grade_id'] is None) == (grade['grading_status'] == 'undetermined'), '空品级状态不符')
            boundary = grade.get('boundary') or item.get('editorial_boundary') or as_text(item.get('editorial_analysis'))
            relations = item.get('related_existing_or_candidates') if namespace == 'html96' else item['related_existing']
            rows.append({'namespace': namespace, 'original_id': item['candidate_id'], 'name': item['name'], 'ontology_type': item['ontology_type'],
                         'grade_id': grade['provisional_grade_id'], 'grading_status': grade['grading_status'], 'entry_status': '未入登记库',
                         'task_family': grade['task_family'], 'original_text': item['original_text'],
                         'boundary': boundary or '此前本选区未单独完善编者边界；不借邻条补写。',
                         'reason': grade['reason'], 'falsifier': grade['falsifier'], 'confidence': grade['confidence'],
                         'content_version': '冻结选区；未设登记版本',
                         'source': as_text({'selection': item['source_selection'], 'declared_upstream': item['declared_upstream'],
                                          'original_proposal_context': {key: grade_document.get(key) for key in ('batch_id', 'policy_id', 'policy_version', 'basis', 'assessor')}}),
                         'content_sha256': item['text_sha256'], 'relations': as_text(relations), 'assessed_at': grade_document['created_at']})
    analyzed = unique(documents['solve24_analysis']['items'], lambda item: item['source_id'])
    raw_items = unique(documents['solve_raw']['items'], lambda item: item['source_id'])
    for item in analyzed.values():
        raw, grade = raw_items[item['source_id']], item['editorial_analysis']
        require(raw['source_version'] == item['source_version'] and raw['source_selection'] == item['source_selection']
                and grade['source_id'] == item['source_id'], 'Solve版本、选区或意见外键不符')
        text = as_text(raw['original_entry'])
        rows.append({'namespace': 'solve24', 'original_id': item['source_id'], 'name': raw['name'], 'ontology_type': grade['ontology_type'],
                     'grade_id': grade['grade_id'], 'grading_status': 'provisional', 'entry_status': '未入登记库',
                     'task_family': grade['task_family'], 'original_text': text,
                     'boundary': as_text({'when': item['source_when'], 'not_when': item['source_not_when'], 'preconditions': item['source_preconditions'],
                                          'editorial_caveat': grade['editorial_caveat'], 'registration_recommendation': grade['registration_recommendation']}),
                     'reason': grade['reason'], 'falsifier': grade['falsifier'], 'confidence': grade['confidence'], 'content_version': item['source_version'],
                     'source': as_text(item['source_selection']), 'content_sha256': digest(text), 'relations': as_text(grade['related_existing']),
                     'assessed_at': documents['solve24_analysis']['assessed_at']})
    names = defaultdict(list)
    for row in rows:
        row['key'] = row['namespace'] + ':' + row['original_id'] + ('@' + row['content_version'] if row['namespace'] == 'registry' else '')
        require(row['ontology_type'] in TYPE_LABELS and (row['grade_id'] is None or row['grade_id'] in GRADES), '未知类型或词表外品级')
        require(row['confidence'] in {'', 'low', 'medium'}, '未知置信标签')
        names[row['name']].append(row['key'])
    unique(rows, lambda row: row['key'])
    for row in rows:
        others = [key for key in names[row['name']] if key != row['key']]
        if others:
            row['relations'] += '\n同名限定内容未自动合并：' + '；'.join(others)
    rank = {grade: i for i, grade in enumerate(GRADES)}
    rows.sort(key=lambda row: rank.get(row['grade_id'], -1), reverse=True)
    return rows


def grade_label(row):
    return GRADE_LABELS[row['grade_id']] if row['grade_id'] else {'undetermined': '暂不可判', 'unrated': '未初评', 'multiple-contexts': '多情境意见'}[row['grading_status']]


def grid(rows):
    result = []
    for number, row in enumerate(rows, 1):
        values = [number, grade_label(row), row['name'], TYPE_LABELS[row['ontology_type']], row['entry_status'], BATCH_LABELS[row['namespace']],
                  row['task_family'], row['original_text'], row['boundary'], row['reason'], row['falsifier'], row['confidence'], row['original_id'],
                  row['content_version'], row['source'], row['content_sha256'], row['relations'], row['assessed_at']]
        for value in values:
            if isinstance(value, str):
                require(len(value.encode('utf-16-le')) // 2 <= 32767, 'Excel单元格超过32767 UTF-16单位，不能静默截断')
                require(not re.search(r'[\x00-\x08\x0b-\x0c\x0e-\x1f\ufffe\uffff]', value), 'Excel原文含不支持的字符，不能静默清洗')
        result.append(values)
    return result


def notes(rows, binding):
    counts = Counter(row['namespace'] for row in rows)
    return ['全部功法总表（保留原初评，不重评）',
            f"合计{len(rows)}条；" + '；'.join(f'{BATCH_LABELS[ns]} {counts[ns]}' for ns in BATCH_LABELS)
            + '；输入：' + digest(binding['registry_sha256'] + binding['candidate_sha256']),
            f"品级是编者前瞻；来源候选不自动注册；当前正式评级记录数：{binding['registry_assessments']}。同名未合并。"]


def render_markdown(rows, binding):
    def cell(value):
        return escape(str(value), quote=False).replace('|', '\\|').replace('\r', ' ').replace('\n', ' ')
    counts = Counter(row['entry_status'] for row in rows)
    provisional = sum(row['grading_status'] == 'provisional' for row in rows)
    undetermined = sum(row['grading_status'] == 'undetermined' for row in rows)
    lines = ['# 功法总表', '', MARKER, '', '## 字多不看', '',
             f"- 共**{len(rows)}条内容记录**；已登记{counts['已登记']}条，未入登记库的来源候选{counts['未入登记库']}条。不是独立功法去重后的总数。",
             f'- 原前瞻初评{provisional}条，暂不可判{undetermined}条；不是正式效果评级，亦不把缺测默认补成黄阶。',
             '- 同一总表的[Excel完整视图](catalog.xlsx)含限定原文/源条目、来源选区、边界、理由、反证和置信。',
             '- 活跃登记来自[registry.json](registry.json)，历史候选来自[仓内快照索引](catalog-sources.json)；维护边界见[README](README.md)。', '',
             '输入绑定：`' + digest(binding['registry_sha256'] + binding['candidate_sha256']) + '`。', '',
             '## 全部记录', '', '| 序号 | 暂定品级 | 功法名称 | 本体类别 | 登记状态 | 来源批次 | 用途/任务族 | 原始ID / 版本 |',
             '|---|---|---|---|---|---|---|---|']
    for number, row in enumerate(rows, 1):
        values = [number, grade_label(row), row['name'], TYPE_LABELS[row['ontology_type']], row['entry_status'],
                  BATCH_LABELS[row['namespace']], row['task_family'], row['original_id'] + ' / ' + row['content_version']]
        lines.append('| ' + ' | '.join(cell(value) for value in values) + ' |')
    lines += ['', '## 更新与核对', '', '```bash', 'make sync-gongfa-catalog', 'make check-gongfa-catalog', 'make test-gongfa-catalog', '```', '',
              '不要手改两份生成视图；候选来源变更另建快照并更新索引，登记内容变更遵守既有版本/初评协议。',
              '内容新版本不继承旧初评；同对象不同情境的活动意见并列保留，不按最高档或最新时间裁决。',
              '本视图不执行方法，不证明原著、许可证、独立语义批准或真实效果；未审材料不混入。', '']
    return '\n'.join(lines)


def build_workbook(rows, binding):
    values = grid(rows)
    book = Workbook()
    book.properties.creator = CREATOR
    book.properties.title = '全部功法总表'
    book.properties.subject = '生成视图；来源候选与登记分开，不是正式评级'
    sheet = book.active
    sheet.title = '功法总表'
    sheet.sheet_view.showGridLines = False
    font = Font(name='Microsoft YaHei', size=10)
    alignment = Alignment(vertical='top', wrap_text=True)
    for number, note in enumerate(notes(rows, binding), 1):
        sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=18)
        item = sheet.cell(number, 1, note)
        item.data_type = 's'
        item.font = Font(name='Microsoft YaHei', size=14 if number == 1 else 10, bold=number == 1)
        item.alignment = alignment
        sheet.row_dimensions[number].height = 27
    sheet.append(HEADERS)
    for item in sheet[4]:
        item.data_type = 's'
        item.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
        item.fill = PatternFill('solid', fgColor='244568')
        item.alignment = alignment
    sheet.row_dimensions[4].height = 32
    for row in values:
        sheet.append(row)
        for item in sheet[sheet.max_row]:
            if isinstance(item.value, str):
                item.data_type = 's'  # 不让来源字符串触发Excel公式解释。
            item.font, item.alignment = font, alignment
            if row[1] == '暂不可判':
                item.fill = PatternFill('solid', fgColor='FFF0CC')
        sheet.row_dimensions[sheet.max_row].height = 40
    for number, width in enumerate([7, 11, 34, 19, 15, 20, 38, 65, 55, 55, 55, 12, 46, 24, 62, 68, 60, 30], 1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.freeze_panes, sheet.auto_filter.ref = 'G5', f'A4:R{sheet.max_row}'
    sheet.print_title_rows = '1:4'
    return book


def read_workbook(path):
    raw = protocol.bounded_read(path, 16 * 1024 * 1024)
    with ZipFile(BytesIO(raw)) as archive:
        members = archive.infolist()
        require(len(members) <= 100 and sum(item.file_size for item in members) <= 64 * 1024 * 1024, 'Excel压缩包超过100项或64MiB展开预算')
        allowed = {'[Content_Types].xml', '_rels/.rels', 'docProps/app.xml', 'docProps/core.xml',
                   'xl/theme/theme1.xml', 'xl/worksheets/sheet1.xml', 'xl/styles.xml', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels'}
        require({item.filename for item in members} == allowed and len(members) == len(allowed),
                'Excel含非生成包成员；不接受宏、外链、连接或额外工作表')
        for item in members:
            xml = archive.read(item)
            require(b'\x00' not in xml and b'<!DOCTYPE' not in xml and b'<!ENTITY' not in xml,
                    'Excel只接受生成的UTF-8 XML，不接受实体或DTD')
            if item.filename.endswith('.rels'):
                relations = ElementTree.fromstring(xml)
                require(all(node.get('TargetMode', '').lower() != 'external' and ':' not in node.get('Target', '')
                            and not node.get('Target', '').startswith(('//', '\\\\')) for node in relations), 'Excel关系不能指向外部资源')
        workbook = ElementTree.fromstring(archive.read('xl/workbook.xml'))
        for node in workbook.findall('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}definedName'):
            valid = node.get('localSheetId') == '0' and (
                node.get('name') == '_xlnm.Print_Titles' and node.text == "'功法总表'!$1:$4" or
                node.get('name') == '_xlnm._FilterDatabase' and re.fullmatch(r"'功法总表'!\$A\$4:\$R\$[0-9]+", node.text or ''))
            require(valid, 'Excel不接受额外命名公式或跨表定义')
    return load_workbook(BytesIO(raw), data_only=False)


def view_path(root, name):
    nominal = root.resolve() / 'metadata/gongfa' / name
    require(not nominal.is_symlink(), '生成目标不能是软链接')
    path = protocol.safe_path(root.resolve(), 'metadata/gongfa/' + name)
    require(path == nominal, '生成目录不能通过软链接重定向')
    return path


def check_views(root, rows, binding):
    markdown = view_path(root, 'catalog.md')
    require(protocol.bounded_read(markdown, 4 * 1024 * 1024).decode('utf-8') == render_markdown(rows, binding), 'catalog.md陈旧或手改；运行make sync-gongfa-catalog')
    book = read_workbook(view_path(root, 'catalog.xlsx'))
    try:
        require(book.properties.creator == CREATOR and book.sheetnames == ['功法总表'], 'Excel不是本工具的单工作表视图')
        sheet = book['功法总表']
        require(sheet.max_row == len(rows) + 4 and sheet.max_column == 18 and sheet.freeze_panes == 'G5'
                and sheet.auto_filter.ref == f'A4:R{len(rows) + 4}', 'Excel范围、筛选或冻结表头不符')
        require([sheet.cell(i, 1).value for i in range(1, 4)] == notes(rows, binding), 'Excel输入绑定或范围说明不符')
        require([item.value for item in sheet[4]] == HEADERS, 'Excel表头不符')
        actual = [[item.value if item.value is not None else '' for item in line] for line in sheet.iter_rows(min_row=5)]
        require(actual == grid(rows), 'Excel正文或初评陈旧/手改；运行make sync-gongfa-catalog')
        require(all(item.data_type != 'f' and item.hyperlink is None for line in sheet.iter_rows() for item in line), 'Excel含公式或可执行外链')
    finally:
        book.close()


def write_views(root, rows, binding):
    markdown = render_markdown(rows, binding).encode('utf-8')
    book = build_workbook(rows, binding)  # 先检查所有文字上限，再触碰旧视图。
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    targets = [(view_path(root, 'catalog.md'), markdown), (view_path(root, 'catalog.xlsx'), buffer.getvalue())]
    before = {}
    for path, _data in targets:
        require(not path.is_symlink(), '生成目标不能是软链接')
        before[path] = sha256(path.read_bytes()).hexdigest() if path.exists() else None
        if path.exists() and path.suffix == '.md':
            require(MARKER in protocol.bounded_read(path, 4 * 1024 * 1024).decode('utf-8')[:200], '拒绝覆盖非本工具生成的Markdown')
        elif path.exists():
            old = read_workbook(path)
            try:
                require(old.properties.creator == CREATOR, '拒绝覆盖非本工具生成的Excel')
            finally:
                old.close()
    for path, data in targets:
        current = sha256(path.read_bytes()).hexdigest() if path.exists() else None
        require(current == before[path], '生成目标被并行修改，请先核对后重跑')
        if path.exists() and path.read_bytes() == data:
            continue
        with tempfile.NamedTemporaryFile(prefix='.' + path.stem + '-', suffix='.tmp', dir=path.parent, delete=False) as handle:
            handle.write(data)
            temporary = Path(handle.name)
        require((sha256(path.read_bytes()).hexdigest() if path.exists() else None) == before[path],
                '提交视图前目标被并行修改；保留临时文件，不覆盖目标')
        # 单文件原子替换；两种呈现之间若被中断，--check会暴露漂移，重跑即可恢复。
        temporary.replace(path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=ROOT)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--write', action='store_true')
    mode.add_argument('--check', action='store_true')
    args = parser.parse_args()
    try:
        root = args.root.resolve(strict=True)
        documents, binding = load_inputs(root)
        rows = merge_rows(documents)
        if args.write:
            write_views(root, rows, binding)
        check_views(root, rows, binding)
        _documents, fresh_binding = load_inputs(root)
        require(fresh_binding == binding, '输入在生成/检查期间变化；视图不宣称最新，请重跑')
        print(json.dumps({'records': len(rows), 'batches': dict(Counter(row['namespace'] for row in rows)),
                          'provisional': sum(row['grading_status'] == 'provisional' for row in rows),
                          'undetermined': sum(row['grading_status'] == 'undetermined' for row in rows),
                          'registered': sum(row['entry_status'] == '已登记' for row in rows),
                          'formal_assessments': binding['registry_assessments'], 'mode': 'write' if args.write else 'check'}, ensure_ascii=False))
        return 0
    except Exception as error:
        # 第三方异常可能包含输入值；只报告已控的位置/规则或异常类别。
        detail = str(error) if type(error) in (ValueError, protocol.ContractError) and not isinstance(error, json.JSONDecodeError) else type(error).__name__
        print('GONGFA_CATALOG_ERRORS\n' + detail, file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
