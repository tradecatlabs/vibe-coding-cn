#!/usr/bin/env python3
"""总表维护的相关行为契约；离线保存隔离工件，不验证功法效果。"""
import argparse
from collections import Counter
from copy import deepcopy
from hashlib import sha256
from html import unescape
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'scripts/sync-gongfa-catalog.py'
ARTIFACTS = None
catalog = None


def save(path, value):
    with path.open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def parse_pipe_row(line):
    cells = []
    current = []
    index = 1
    while index < len(line) - 1:
        if line[index:index + 2] == r'\|':
            current.append('|')
            index += 2
        elif line[index] == '|':
            cells.append(''.join(current).strip())
            current = []
            index += 1
        else:
            current.append(line[index])
            index += 1
    cells.append(''.join(current).strip())
    return cells


class CatalogMaintenanceTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.documents, cls.binding = catalog.load_inputs(ROOT)

    def isolated_project(self):
        root = Path(tempfile.mkdtemp(prefix='project-', dir=ARTIFACTS))
        target = root / 'metadata/gongfa'
        target.mkdir(parents=True)
        # 只复制本消费链所需输入；不依赖原HTML、完整Solve或个人目录。
        for path in (ROOT / 'metadata/gongfa').iterdir():
            if path.name in {'registry.json', 'registry.schema.json', 'README.md', 'catalog-sources.json'} or path.name.startswith('source-'):
                shutil.copyfile(path, target / path.name)
        return root

    def command(self, root, *options):
        argv = [sys.executable, str(SCRIPT), '--root', str(root), *options]
        result = subprocess.run(argv, cwd=ARTIFACTS, capture_output=True, text=True, timeout=60)
        log = Path(tempfile.mkdtemp(prefix='command-', dir=ARTIFACTS))
        save(log / 'result.json', {'argv': argv, 'cwd': str(ARTIFACTS), 'exit_code': result.returncode, 'timeout_seconds': 60})
        for name, text in [('stdout.log', result.stdout), ('stderr.log', result.stderr)]:
            with (log / name).open('x', encoding='utf-8') as stream:
                stream.write(text)
        return result

    def test_current_all_classified_content_and_two_bad_summaries_keep_original_opinions(self):
        rows = catalog.merge_rows(deepcopy(self.documents))
        self.assertEqual(Counter(row['namespace'] for row in rows), {'registry': 79, 'html211': 211, 'solve24': 24, 'html96': 96})
        self.assertEqual(len(rows), 410)
        self.assertEqual(len({(r['namespace'], r['original_id'], r['content_version']) for r in rows}), 410)
        undetermined = {(r['namespace'], r['original_id']) for r in rows if r['grading_status'] == 'undetermined'}
        self.assertEqual(undetermined, {('html211', 'E001'), ('html96', 'VC029')})
        by_id = {(r['namespace'], r['original_id']): r for r in rows}
        self.assertEqual(by_id['registry', 'gongfa-first-principles']['grade_id'], 'di-advanced')
        self.assertEqual(by_id['solve24', 'psoa.programming.minimal-reproduction']['grade_id'], 'di-initial')
        self.assertIn('一个分布式系统最多只能保证一致性、可用性和分区容错性中的两项。', by_id['html96', 'VC029']['original_text'])
        self.assertIn('失去50%容量时剩余单位需要承担2倍总流量。', by_id['html211', 'E001']['original_text'])
        self.assertEqual(sum(r['entry_status'] == '已登记' for r in rows), 79)
        self.assertEqual(self.documents['registry']['assessments'], [])
        registry_item = next(item for item in self.documents['registry']['gongfa'] if item['id'] == 'gongfa-first-principles')
        profile = registry_item['content']['profile']
        profile_parts = [profile['summary'], '问题：' + profile['problem'], '机制：' + profile['mechanism']]
        for key, label in (('conditions', '条件'), ('limitations', '限制'), ('expected_outcomes', '预期结果')):
            if profile[key]:
                profile_parts.append(label + '：' + '；'.join(profile[key]))
        self.assertEqual(' '.join(profile_parts), by_id['registry', 'gongfa-first-principles']['explanation'])
        html_grade = next(item for item in self.documents['html211_grades']['items'] if item['candidate_id'] == 'E001')
        html_candidate = next(item for item in self.documents['html211_candidates']['candidates'] if item['candidate_id'] == 'E001')
        expected_html_explanation = (html_candidate['original_text'] + ' 用途/任务族：' + html_grade['task_family'] +
                                     ' 初评说明：' + html_grade['reason'])
        self.assertEqual(expected_html_explanation, by_id['html211', 'E001']['explanation'])
        solve_item = self.documents['solve24_analysis']['items'][0]
        solve_grade = solve_item['editorial_analysis']
        solve_row = by_id['solve24', solve_grade['source_id']]
        expected_solve_explanation = '；'.join((
            '用途/任务族：' + solve_grade['task_family'],
            '机制：' + solve_grade['mechanism'],
            '适用：' + '；'.join(solve_item['source_when']),
            '不适用：' + '；'.join(solve_item['source_not_when']),
            '前提：' + '；'.join(solve_item['source_preconditions']),
            '边界：' + solve_grade['editorial_caveat'],
        ))
        self.assertEqual(expected_solve_explanation, solve_row['explanation'])
        markdown = catalog.render_markdown(rows, self.binding)
        self.assertNotIn('[控制论与科学方法论](#控制论与科学方法论)', markdown)
        section = markdown.split('## 全部记录\n\n', 1)[1].split('\n\n## 更新与核对', 1)[0]
        table_rows = [line for line in section.splitlines() if line.startswith('|')]
        cells = [parse_pipe_row(line) for line in table_rows]
        self.assertEqual(['品级', '功法名称', '解释'], cells[0])
        self.assertEqual(len(rows) + 2, len(cells))
        self.assertTrue(all(len(row) == 3 for row in cells), 'Markdown浏览表只保留品级、名称和解释')
        self.assertEqual([row['name'] for row in rows], [row[1] for row in cells[2:]])
        expected_explanations = [row['explanation'].replace('\r', ' ').replace('\n', ' ') for row in rows]
        self.assertEqual(expected_explanations, [unescape(row[2]) for row in cells[2:]])

    def test_registry_updates_follow_explicit_correction_and_never_inherit_a_new_version(self):
        documents = deepcopy(self.documents)
        batch = documents['registry']['grade_proposals'][0]
        correction = deepcopy(batch)
        correction['id'] = 'test-explicit-correction'
        correction['supersedes'] = batch['id']
        correction['assessed_at'] = '2026-10-08T00:00:00Z'
        for item in correction['items']:
            if item['gongfa_id'] == 'gongfa-first-principles':
                item['grade_id'] = 'xuan-intermediate'
                item['reason'] = '隔离夹具的显式更正。'
        documents['registry']['grade_proposals'].append(correction)
        new_version = deepcopy(next(item for item in documents['registry']['gongfa'] if item['id'] == 'gongfa-first-principles'))
        new_version['content_version'] = 'r2'
        new_version['predecessor_version'] = 'r1'
        new_version['content']['original_text'] = '新版本的限定内容尚无初评。'
        new_version['content']['text_sha256'] = sha256('新版本的限定内容尚无初评。'.encode('utf-8')).hexdigest()
        documents['registry']['gongfa'].append(new_version)
        selected = {row['content_version']: row for row in catalog.merge_rows(documents) if row['namespace'] == 'registry' and row['original_id'] == 'gongfa-first-principles'}
        self.assertEqual(selected['r1']['grade_id'], 'xuan-intermediate')
        self.assertEqual(selected['r1']['reason'], '隔离夹具的显式更正。')
        self.assertIsNone(selected['r2']['grade_id'])
        self.assertEqual(selected['r2']['grading_status'], 'unrated')
        another = deepcopy(correction)
        another['id'] = 'test-another-context'
        another['supersedes'] = None
        another['application_context'] = '隔离夹具的不同任务情境'
        documents['registry']['grade_proposals'].append(another)
        row = next(r for r in catalog.merge_rows(documents) if r['namespace'] == 'registry' and r['original_id'] == 'gongfa-first-principles' and r['content_version'] == 'r1')
        self.assertEqual(row['grading_status'], 'multiple-contexts')
        self.assertIsNone(row['grade_id'])
        self.assertIn('test-another-context', row['reason'])

    def test_portable_inputs_reject_drift_and_out_of_project_paths(self):
        root = self.isolated_project()
        result = self.command(root, '--write')
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(self.command(root, '--check').returncode, 0)
        source_index = json.loads((root / 'metadata/gongfa/catalog-sources.json').read_text())
        snapshot = root / 'metadata/gongfa' / source_index['snapshot']['path']
        with snapshot.open('ab') as stream:
            stream.write(b'\n')
        self.assertEqual(self.command(root, '--check').returncode, 1)
        index = root / 'metadata/gongfa/catalog-sources.json'
        source_index['snapshot']['path'] = '../outside.json'
        # 这是本测试独占的隔离夹具，不修改项目索引。
        index.write_text(json.dumps(source_index), encoding='utf-8')
        self.assertEqual(self.command(root, '--check').returncode, 1)

    def test_excel_text_safety_stale_detection_and_generated_file_ownership(self):
        documents = deepcopy(self.documents)
        item = documents['html96_candidates']['candidates'][0]
        grade = documents['html96_grades']['items'][0]
        item['name'] = grade['name'] = '=2+2'
        rows = catalog.merge_rows(documents)
        root = Path(tempfile.mkdtemp(prefix='excel-', dir=ARTIFACTS))
        target = root / 'metadata/gongfa'
        target.mkdir(parents=True)
        catalog.write_views(root, rows, self.binding)
        catalog.check_views(root, rows, self.binding)
        book = load_workbook(target / 'catalog.xlsx', data_only=False)
        self.assertEqual(book.sheetnames, ['功法总表'])
        sheet = book['功法总表']
        self.assertEqual(sheet.freeze_panes, 'G5')
        self.assertEqual(sheet.auto_filter.ref, 'A4:R414')
        self.assertIn('=2+2', [sheet.cell(i, 3).value for i in range(5, 415)])
        for row in sheet.iter_rows():
            for cell in row:
                self.assertNotEqual(cell.data_type, 'f')
                self.assertIsNone(cell.hyperlink)
        sheet['B5'] = '黄初'
        book.save(target / 'catalog.xlsx')
        book.close()
        with self.assertRaises(ValueError):
            catalog.check_views(root, rows, self.binding)
        for text in ['😀' * 16384, '非法\x01字符']:
            bad = deepcopy(rows)
            bad[0]['name'] = text
            with self.assertRaises(ValueError):
                catalog.write_views(root, bad, self.binding)
        hand_written = Path(tempfile.mkdtemp(prefix='owned-', dir=ARTIFACTS))
        own_target = hand_written / 'metadata/gongfa'
        own_target.mkdir(parents=True)
        (own_target / 'catalog.md').write_text('手工内容，不能覆盖。', encoding='utf-8')
        with self.assertRaises(ValueError):
            catalog.write_views(hand_written, rows, self.binding)
        self.assertEqual((own_target / 'catalog.md').read_text(), '手工内容，不能覆盖。')


def main():
    global ARTIFACTS, catalog
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--artifacts', type=Path)
    args = parser.parse_args()
    ARTIFACTS = args.artifacts or Path(tempfile.mkdtemp(prefix='gongfa-catalog-test-'))
    if args.artifacts:
        ARTIFACTS.mkdir(parents=True, exist_ok=False)
    ARTIFACTS = ARTIFACTS.resolve()
    spec = importlib.util.spec_from_file_location('gongfa_catalog', SCRIPT)
    catalog = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(catalog)
    result = unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromTestCase(CatalogMaintenanceTest))
    save(ARTIFACTS / 'result.json', {'tests_run': result.testsRun, 'failures': len(result.failures), 'errors': len(result.errors),
                                   'success': result.wasSuccessful(), 'script_sha256': sha256(SCRIPT.read_bytes()).hexdigest()})
    print('测试工件：' + str(ARTIFACTS))
    return 0 if result.wasSuccessful() else 1


if __name__ == '__main__':
    sys.exit(main())
